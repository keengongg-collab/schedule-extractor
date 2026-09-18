# ============================================================
# 轻量化AI文档排班日程提取工具
# Copyright (c) 2026 schedule-extractor Contributors
# Licensed under MIT License
# 详见 LICENSE 文件与 PRIVACY.md 隐私声明
# ============================================================

"""
文档解析器模块
负责读取 PDF、Word(docx)、纯文本文件，提取纯文本内容供 AI 分析

v1.0 支持格式：
- PDF：使用 pdfplumber 读取文本
- DOCX：使用 python-docx 读取段落和表格（不支持旧版 .doc）
- 纯文本：直接读取
"""
import os


def parse_document(filepath: str) -> str:
    """
    根据文件类型解析文档，返回纯文本内容
    :param filepath: 文件路径
    :return: 提取的纯文本
    """
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        return _parse_pdf(filepath)
    elif ext == ".docx":
        return _parse_word(filepath)
    elif ext == ".txt":
        return _parse_text(filepath)
    else:
        raise ValueError(f"不支持的文件类型: {ext}（v1.0 仅支持 PDF / DOCX / TXT）")


def _parse_pdf(filepath: str) -> str:
    """读取 PDF 文件内容"""
    import pdfplumber

    text_parts = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            # 提取普通文本
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

            # 提取表格文本（排班表常以表格形式存在）
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    # 过滤空行，拼接单元格
                    cells = [cell for cell in row if cell and cell.strip()]
                    if cells:
                        text_parts.append(" | ".join(cells))

    return "\n".join(text_parts)


def _parse_word(filepath: str) -> str:
    """读取 Word(.docx) 文件内容（段落 + 表格）

    注意：python-docx 仅支持 OOXML 格式的 .docx，不支持旧版二进制 .doc。
    """
    from docx import Document

    doc = Document(filepath)
    text_parts = []

    # 读取段落
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)

    # 读取表格（排班表常见格式）
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                text_parts.append(" | ".join(cells))

    return "\n".join(text_parts)


def _parse_text(filepath: str) -> str:
    """读取纯文本文件"""
    # 尝试多种编码，避免乱码
    for encoding in ["utf-8", "gbk", "gb2312", "ascii"]:
        try:
            with open(filepath, "r", encoding=encoding) as f:
                return f.read()
        except (UnicodeDecodeError, FileNotFoundError):
            continue
    raise ValueError("无法读取文本文件，编码不支持")


def parse_text(text: str) -> str:
    """
    直接处理纯文本（非文件来源）
    用于文本粘贴场景，做简单的清洗
    """
    # 去除多余空行
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)
